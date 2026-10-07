from pathlib import Path

import openpyxl

from parser.pdf.type import PageText
from parser.xlsx.type import TableGrid, TableCell


def _fill_merged_cells(ws, grid: list[list[str]]) -> None:
    """병합 셀 범위를 좌상단 값으로 채워서 NaN처럼 비어 보이지 않게 한다."""
    for merged_range in ws.merged_cells.ranges:
        min_row, min_col, max_row, max_col = (
            merged_range.min_row,
            merged_range.min_col,
            merged_range.max_row,
            merged_range.max_col,
        )
        value = grid[min_row - 1][min_col - 1]
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                grid[r - 1][c - 1] = value


def _build_table_grid(title:str, ws) -> TableGrid:
    cells = []
    covered = set()

    merge_info = {}
    for mr in ws.merged_cells.ranges:
        merge_info[(mr.min_row, mr.min_col)] = (
            mr.max_row - mr.min_row + 1,
            mr.max_col - mr.min_col + 1,
        )
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                if (r, c) != (mr.min_row, mr.min_col):
                    covered.add((r - 1, c - 1))

    for row in ws.iter_rows():
        for cell in row:
            r, c = cell.row, cell.column
            if (r - 1, c - 1) in covered:
                continue
            rowspan, colspan = merge_info.get((r, c),(1, 1))
            cells.append(
                TableCell(
                    row = r - 1,
                    col = c - 1,
                    text= "" if cell.value is None else str(cell.value),
                    rowspan=rowspan,
                    colspan=colspan,
            ))
    return TableGrid(
        title=title,
        table_cells=cells,
        covered=covered)


def _grid_to_html_table(table_grid:TableGrid):
    if not table_grid.table_cells:
        return "<table></table>"

    table_cells = table_grid.table_cells

    row_max = max(cell.row + cell.rowspan for cell in table_cells)
    col_max = max(cell.col + cell.colspan for cell in table_cells)
    table = [[None for _ in range(col_max)] for _ in range(row_max)]

    for cell in table_cells:
        table[cell.row][cell.col] = cell

    covered_cells = table_grid.covered
    tr_list = []

    for r in range(row_max):
        td_list = []
        for c in range(col_max):
            if (r, c) in covered_cells:
                continue
            cell = table[r][c]
            if not cell:
                td_list.append("<td></td>")
                continue

            row_span = getattr(cell,"rowspan",1) or 1
            col_span = getattr(cell,"colspan",1) or 1

            # 속성 태그 문자열 구성
            attrs = ""
            if row_span > 1:
                attrs += f' rowspan="{row_span}"'
            if col_span > 1:
                attrs += f' colspan="{col_span}"'

            td_list.append(f"<td{attrs}>{cell.text}</td>")
        tr = "  <tr>\n    " + "\n    ".join(td_list) + "\n  </tr>"
        tr_list.append(tr)
    return "<table>\n" + "\n".join(tr_list) + "\n</table>"

def parser_xlsx_md(
        file_path:Path,
) -> str:
    return "\n\n".join(page.text for page in parser_xlsx_pages(file_path))


def parser_xlsx_pages(file_path: Path) -> list[PageText]:
    """시트 하나를 한 페이지로 본다. page에는 시트 순번(1부터)이 들어간다."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    pages: list[PageText] = []
    try:
        for i, ws in enumerate(wb.worksheets, start=1):
            table_grid = _build_table_grid(ws.title, ws)
            table = _grid_to_html_table(table_grid)
            pages.append(PageText(page=i, text=f"## {ws.title}\n\n{table}"))
    finally:
        wb.close()
    return pages
