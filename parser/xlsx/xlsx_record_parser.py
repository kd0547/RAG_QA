"""XLSX를 레코드(행) 단위 텍스트로 변환하는 파서.

parser_xlsx_pages(HTML 표)는 시트 전체를 한 덩어리로 넘겨서 청커가 셀 중간을 자르고,
헤더 없이 숫자만 남은 청크가 생긴다. 이 파서는 행 하나를 "헤더: 값 | ..." 한 줄로
직렬화하고, chunk_size 이하로 미리 묶어 PageText로 반환한다. 그래서 청커의 고정 간격
분할이 레코드 중간을 자르지 않는다.

행 단위로만 처리해도 전치형 표(행=속성, 열=제품)가 깨지지 않는다. 각 줄에
속성명(첫 열)과 제품명(헤더), 값이 함께 들어가기 때문이다.
"""
from __future__ import annotations

import datetime
import logging
from pathlib import Path

import openpyxl

from config import DEFAULT_CHUNK_SIZE
from parser.pdf.type import PageText
from parser.xlsx.xlsx_parser import _fill_merged_cells

logger = logging.getLogger(__name__)

# 수식 오류값은 의미가 없으므로 빈 값처럼 생략한다. (TBD 같은 일반 문자열은 남긴다)
_ERROR_VALUES = frozenset({"#N/A", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!"})


def _cell_text(value) -> str:
    """셀 값을 한 줄 문자열로 바꾼다."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime):
        return value.date().isoformat() if value.time() == datetime.time() else value.isoformat(sep=" ")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()

    text = " ".join(str(value).split())  # 셀 안 줄바꿈 때문에 레코드가 여러 줄로 갈라지지 않게 한다
    return "" if text in _ERROR_VALUES else text


def _build_grid(ws) -> list[list[str]]:
    """병합 셀을 좌상단 값으로 채운 2차원 그리드를 만든다."""
    grid = [
        [_cell_text(cell.value) for cell in row]
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)
    ]
    _fill_merged_cells(ws, grid)
    return grid


def _horizontal_merge_rows(ws) -> set[int]:
    """가로로 병합된 셀이 걸쳐 있는 행(0부터)."""
    return {
        r - 1
        for mr in ws.merged_cells.ranges
        if mr.max_col > mr.min_col
        for r in range(mr.min_row, mr.max_row + 1)
    }


def _split_blocks(grid: list[list[str]]) -> list[list[int]]:
    """완전히 빈 행을 경계로 표 블록을 나눈다. 블록은 행 번호(0부터) 목록이다."""
    blocks, current = [], []
    for r, row in enumerate(grid):
        if any(row):
            current.append(r)
        elif current:
            blocks.append(current)
            current = []
    if current:
        blocks.append(current)
    return blocks


def _distinct_values(row: list[str]) -> list[str]:
    seen = []
    for v in row:
        if v and v not in seen:
            seen.append(v)
    return seen


def _compose_header(parts: list[str]) -> str:
    """다단 헤더를 위에서 아래로 합친다. 세로 병합으로 반복된 값과 빈 칸은 뺀다."""
    out = []
    for p in parts:
        if p and (not out or out[-1] != p):
            out.append(p)
    return " / ".join(out)


def _serialize_block(
    grid: list[list[str]],
    rows: list[int],
    merge_rows: set[int],
    prefix_base: str,
    title: str | None,
) -> list[tuple[str, list[str]]]:
    """표 블록 하나를 (접두어, 필드 목록) 레코드 목록으로 바꾼다."""
    # 블록 안에서 값이 하나도 없는 열은 버린다.
    cols = [c for c in range(len(grid[0])) if any(grid[r][c] for r in rows)]

    # 첫 행에 값이 하나뿐이고(병합 제목 포함) 열이 여러 개면 표 제목으로 본다.
    if len(rows) > 1 and len(cols) > 1 and len(_distinct_values([grid[rows[0]][c] for c in cols])) == 1:
        title = _distinct_values(grid[rows[0]])[0]
        rows = rows[1:]

    # 헤더: 상단의 가로 병합 행들 + 그 바로 아래 행(병합 범위를 세분하는 하위 헤더).
    # 병합이 없으면 첫 행만 헤더로 쓴다. 데이터 행은 최소 1개 남긴다.
    merged_top = 0
    while merged_top < len(rows) and rows[merged_top] in merge_rows:
        merged_top += 1
    header_count = min(merged_top + 1, len(rows) - 1)

    header_rows, data_rows = rows[:header_count], rows[header_count:]
    headers = [_compose_header([grid[r][c] for r in header_rows]) for c in cols]

    prefix = f"[{prefix_base} / {title}]" if title else f"[{prefix_base}]"
    logger.info(
        "xlsx 표 판별: %s rows=%d~%d header_rows=%d data_rows=%d cols=%d",
        prefix, rows[0] + 1, rows[-1] + 1, len(header_rows), len(data_rows), len(cols),
    )

    records = []
    for r in data_rows:
        fields = []
        for header, c in zip(headers, cols):
            value = grid[r][c]
            if not value or value == header:  # 빈 값, 헤더까지 이어진 세로 병합 값은 생략
                continue
            fields.append(f"{header}: {value}" if header else value)
        if fields:
            records.append((prefix, fields))
    return records


def _record_lines(prefix: str, fields: list[str], size: int) -> list[str]:
    """레코드 하나를 줄로 만든다. size를 넘으면 필드 경계에서 나누고 첫 필드를 반복한다."""
    line = f"{prefix} " + " | ".join(fields)
    if len(line) <= size or len(fields) == 1:
        return [line]

    lead = f"{prefix} {fields[0]}"
    lines, current = [], lead
    for field in fields[1:]:
        candidate = f"{current} | {field}"
        if len(candidate) > size and current != lead:
            lines.append(current)
            current = f"{lead} | {field}"
        else:
            current = candidate
    lines.append(current)
    return lines


def _pack(lines: list[str], size: int) -> list[str]:
    """줄을 size 이하 묶음으로 합친다. 줄 중간에서는 자르지 않는다."""
    packs, current = [], ""
    for line in lines:
        if current and len(current) + 1 + len(line) > size:
            packs.append(current)
            current = line
        else:
            current = f"{current}\n{line}" if current else line
    if current:
        packs.append(current)
    return packs


def parser_xlsx_record_pages(file_path: Path, chunk_size: int = DEFAULT_CHUNK_SIZE) -> list[PageText]:
    """시트별 레코드 묶음을 PageText로 반환한다. page에는 시트 순번(1부터)이 들어간다.

    한 시트가 chunk_size를 넘으면 같은 page 번호의 PageText가 여러 개 나온다.
    """
    file_path = Path(file_path)
    wb = openpyxl.load_workbook(file_path, data_only=True)
    pages: list[PageText] = []
    try:
        for i, ws in enumerate(wb.worksheets, start=1):
            grid = _build_grid(ws)
            if not grid or not grid[0]:
                continue
            merge_rows = _horizontal_merge_rows(ws)
            prefix_base = f"{file_path.name} / {ws.title}"

            lines: list[str] = []
            pending_title: str | None = None
            for rows in _split_blocks(grid):
                values = _distinct_values([v for r in rows for v in grid[r]])
                # 값 하나짜리 행만 있는 블록은 다음 표의 제목으로 넘긴다.
                if len(rows) == 1 and len(values) == 1:
                    pending_title = values[0]
                    continue
                for prefix, fields in _serialize_block(grid, rows, merge_rows, prefix_base, pending_title):
                    lines.extend(_record_lines(prefix, fields, chunk_size))
                pending_title = None

            pages.extend(PageText(page=i, text=text) for text in _pack(lines, chunk_size))
    finally:
        wb.close()
    return pages
